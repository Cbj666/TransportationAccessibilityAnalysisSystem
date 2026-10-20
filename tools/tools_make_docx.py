from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "项目说明书-v3.0.md"
TARGET = ROOT / "武汉个性化等时圈与区位评估平台项目说明书-v3.0.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text.strip())
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)


def add_table(document, rows):
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            value = row[column_index] if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            set_cell_text(cell, value, bold=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
    document.add_paragraph()


def add_code(document, lines):
    paragraph = document.add_paragraph(style="No Spacing")
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(55, 55, 55)
        if index < len(lines) - 1:
            run.add_break()


def add_body(document, text):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Inches(0.28)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def add_list(document, text, ordered=False):
    style = "List Number" if ordered else "List Bullet"
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for name, size, color in [("Title", 22, "1F4E79"), ("Heading 1", 16, "1F4E79"), ("Heading 2", 13, "2F75B5"), ("Heading 3", 11, "2F75B5")]:
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    if "Code" not in [style.name for style in document.styles]:
        code_style = document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        code_style.font.size = Pt(9)


def convert():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    in_code = False
    code_lines = []
    index = 0

    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                add_code(document, code_lines)
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                current = lines[index]
                if not re.match(r"^\|\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$", current):
                    table_lines.append([part.strip() for part in current.strip().strip("|").split("|")])
                index += 1
            if table_lines:
                add_table(document, table_lines)
            continue

        match = re.match(r"^(#{1,3})\s+(.*)$", line)
        if match:
            level = len(match.group(1))
            paragraph = document.add_paragraph(match.group(2), style="Title" if level == 1 else f"Heading {level}")
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue

        match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if match:
            add_list(document, match.group(2), ordered=True)
            index += 1
            continue

        match = re.match(r"^-\s+(.*)$", line)
        if match:
            add_list(document, match.group(1))
            index += 1
            continue

        add_body(document, line)
        index += 1

    document.core_properties.title = "武汉个性化等时圈与区位评估平台项目说明书"
    document.core_properties.subject = "GIS 地理信息系统原理与开发课程项目"
    document.core_properties.author = "项目组"
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    convert()
